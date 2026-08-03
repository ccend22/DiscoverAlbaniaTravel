#!/usr/bin/env python3
"""Scrape public data exposed by https://www.etransport.al/.

The site is JavaScript-driven, so this script uses a real Chromium browser via
Playwright. It captures JSON/API responses, crawls public pages, extracts visible
contacts/prices/company names, and exports XLSX/CSV/JSON files.

No login, authentication bypass, or private endpoints are used.
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import json
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urljoin, urlparse, urldefrag

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.async_api import BrowserContext, Page, Response, async_playwright

BASE_URL = "https://www.etransport.al"
START_PATHS = [
    "/",
    "/RrjetiLinjave/Company/Operatoret",
    "/RrjetiLinjave/Company/Agjencite",
    "/RrjetiLinjave/LinjaDheOrare",
    "/Information/Kontakte",
]

PHONE_RE = re.compile(r"(?:\+?355|00355|0)?[\s()./-]*(?:6\d|4\d|8\d)[\d\s()./-]{6,13}")
EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PRICE_RE = re.compile(
    r"(?:(?:ALL|LEK|LEKË|EUR|EURO|€|\$)\s*\d[\d.,]*|\d[\d.,]*\s*(?:ALL|LEK|LEKË|EUR|EURO|€|\$))",
    re.I,
)
TIME_RE = re.compile(r"\b(?:[01]?\d|2[0-3]):[0-5]\d\b")
URL_RE = re.compile(r"https?://[^\s\"'<>]+", re.I)
API_HINT_RE = re.compile(r"(?:/api/|graphql|operator|agency|agjenci|station|terminal|route|linj|orar|price|cmim)", re.I)

CLICK_TEXTS = [
    "Të tjera",
    "Te tjera",
    "Shfaq më shumë",
    "Shfaq me shume",
    "Më shumë",
    "Me shume",
    "Load more",
    "Next",
    "Pasardhës",
    "Kërko",
    "Kerko",
]

BLOCK_SELECTORS = [
    "article",
    "li",
    "tr",
    "[class*='card']",
    "[class*='item']",
    "[class*='operator']",
    "[class*='company']",
    "[class*='route']",
    "[class*='line']",
    "[class*='station']",
]


@dataclass
class State:
    output_dir: Path
    responses: list[dict[str, Any]] = field(default_factory=list)
    response_hashes: set[str] = field(default_factory=set)
    json_payloads: list[dict[str, Any]] = field(default_factory=list)
    js_urls: set[str] = field(default_factory=set)
    visited: list[dict[str, Any]] = field(default_factory=list)
    blocks: list[dict[str, Any]] = field(default_factory=list)
    links: set[str] = field(default_factory=set)
    errors: list[str] = field(default_factory=list)


def clean_text(value: Any) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"\s+", " ", text).strip()


def normalize_phone(value: str) -> str:
    raw = re.sub(r"[^\d+]", "", value)
    if raw.startswith("00355"):
        raw = "+355" + raw[5:]
    elif raw.startswith("355"):
        raw = "+" + raw
    return raw


def safe_filename(value: str, max_len: int = 110) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("_")
    return value[:max_len] or "file"


def same_origin(url: str) -> bool:
    try:
        parsed = urlparse(url)
        return parsed.scheme in {"http", "https"} and parsed.netloc.lower() in {
            "etransport.al",
            "www.etransport.al",
        }
    except ValueError:
        return False


def canonical_url(url: str) -> str:
    url = urldefrag(url)[0]
    parsed = urlparse(url)
    if not parsed.scheme:
        url = urljoin(BASE_URL, url)
    return url.rstrip("/") or BASE_URL


def looks_like_json(content_type: str, url: str) -> bool:
    ctype = (content_type or "").lower()
    return "json" in ctype or "/api/" in url.lower() or "graphql" in url.lower()


async def capture_response(response: Response, state: State) -> None:
    try:
        url = response.url
        headers = await response.all_headers()
        content_type = headers.get("content-type", "")
        resource_type = response.request.resource_type

        if resource_type == "script" and url.lower().endswith((".js", ".mjs")):
            state.js_urls.add(url)

        if not looks_like_json(content_type, url):
            return

        body = await response.body()
        if not body:
            return
        digest = hashlib.sha256(body).hexdigest()
        if digest in state.response_hashes:
            return
        state.response_hashes.add(digest)

        meta = {
            "url": url,
            "status": response.status,
            "method": response.request.method,
            "resource_type": resource_type,
            "content_type": content_type,
            "sha256": digest,
            "size_bytes": len(body),
        }
        state.responses.append(meta)

        raw_dir = state.output_dir / "raw_json"
        raw_dir.mkdir(parents=True, exist_ok=True)
        path = raw_dir / f"{len(state.responses):04d}_{safe_filename(urlparse(url).path)}_{digest[:10]}.json"
        path.write_bytes(body)

        try:
            payload = json.loads(body.decode("utf-8", errors="replace"))
            state.json_payloads.append({"source_url": url, "payload": payload})
        except json.JSONDecodeError:
            pass
    except Exception as exc:  # noqa: BLE001
        state.errors.append(f"capture_response: {response.url}: {exc}")


async def auto_interact(page: Page, max_rounds: int = 12) -> None:
    """Scroll and click public load-more/pagination/search controls."""
    last_height = 0
    stable_rounds = 0
    for _ in range(max_rounds):
        try:
            await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
            await page.wait_for_timeout(1000)
            height = await page.evaluate("document.body.scrollHeight")
            stable_rounds = stable_rounds + 1 if height == last_height else 0
            last_height = height

            clicked = False
            for label in CLICK_TEXTS:
                loc = page.get_by_text(label, exact=True)
                count = await loc.count()
                for idx in range(min(count, 5)):
                    item = loc.nth(idx)
                    try:
                        if await item.is_visible() and await item.is_enabled():
                            await item.click(timeout=2500)
                            await page.wait_for_timeout(1200)
                            clicked = True
                    except Exception:
                        continue

            # Generic pagination buttons with page numbers or arrows.
            for selector in [
                "button[aria-label*='next' i]",
                "a[aria-label*='next' i]",
                "button:has-text('›')",
                "button:has-text('>')",
            ]:
                loc = page.locator(selector)
                for idx in range(min(await loc.count(), 3)):
                    try:
                        item = loc.nth(idx)
                        if await item.is_visible() and await item.is_enabled():
                            await item.click(timeout=2000)
                            await page.wait_for_timeout(1000)
                            clicked = True
                    except Exception:
                        continue

            if stable_rounds >= 2 and not clicked:
                break
        except Exception:
            break


async def extract_page(page: Page, state: State, url: str) -> None:
    try:
        title = await page.title()
        body_text = clean_text(await page.locator("body").inner_text(timeout=10000))
        state.visited.append(
            {
                "url": page.url,
                "requested_url": url,
                "title": title,
                "body_chars": len(body_text),
                "captured_at": datetime.now(timezone.utc).isoformat(),
            }
        )

        pages_dir = state.output_dir / "pages"
        pages_dir.mkdir(parents=True, exist_ok=True)
        slug = safe_filename(urlparse(page.url).path or "home")
        (pages_dir / f"{slug}.txt").write_text(body_text, encoding="utf-8")
        (pages_dir / f"{slug}.html").write_text(await page.content(), encoding="utf-8")
        try:
            await page.screenshot(path=str(pages_dir / f"{slug}.png"), full_page=True)
        except Exception:
            pass

        hrefs = await page.locator("a[href]").evaluate_all(
            "els => els.map(a => a.href).filter(Boolean)"
        )
        for href in hrefs:
            href = canonical_url(href)
            if same_origin(href):
                state.links.add(href)

        seen_texts: set[str] = set()
        for selector in BLOCK_SELECTORS:
            loc = page.locator(selector)
            count = min(await loc.count(), 3000)
            for idx in range(count):
                try:
                    text = clean_text(await loc.nth(idx).inner_text(timeout=800))
                except Exception:
                    continue
                if len(text) < 4 or text in seen_texts:
                    continue
                seen_texts.add(text)
                phones = sorted({normalize_phone(x) for x in PHONE_RE.findall(text) if len(normalize_phone(x)) >= 7})
                emails = sorted(set(EMAIL_RE.findall(text)))
                prices = sorted(set(PRICE_RE.findall(text)))
                times = sorted(set(TIME_RE.findall(text)))
                if phones or emails or prices or times or any(k in text.lower() for k in ["sh.p.k", "operator", "linja", "nisje", "mbërrit", "stacion"]):
                    state.blocks.append(
                        {
                            "page_url": page.url,
                            "selector": selector,
                            "text": text,
                            "phones": " | ".join(phones),
                            "emails": " | ".join(emails),
                            "prices": " | ".join(prices),
                            "times": " | ".join(times),
                        }
                    )
    except Exception as exc:  # noqa: BLE001
        state.errors.append(f"extract_page: {url}: {exc}")


async def visit(context: BrowserContext, url: str, state: State) -> None:
    page = await context.new_page()
    page.on("response", lambda response: asyncio.create_task(capture_response(response, state)))
    try:
        await page.goto(url, wait_until="domcontentloaded", timeout=60000)
        await page.wait_for_timeout(3500)
        await auto_interact(page)
        await page.wait_for_timeout(1500)
        await extract_page(page, state, url)
    except Exception as exc:  # noqa: BLE001
        state.errors.append(f"visit: {url}: {exc}")
    finally:
        await page.close()


async def inspect_javascript(context: BrowserContext, state: State) -> None:
    js_dir = state.output_dir / "javascript"
    js_dir.mkdir(parents=True, exist_ok=True)
    findings: list[dict[str, str]] = []

    for idx, url in enumerate(sorted(state.js_urls), start=1):
        try:
            response = await context.request.get(url, timeout=30000)
            if not response.ok:
                continue
            text = await response.text()
            if len(text) > 8_000_000:
                continue
            name = f"{idx:04d}_{safe_filename(urlparse(url).path)}.js"
            (js_dir / name).write_text(text, encoding="utf-8", errors="replace")

            candidates = set(URL_RE.findall(text))
            candidates.update(re.findall(r"[\"'](/[^\"']{3,180})[\"']", text))
            for candidate in candidates:
                candidate = clean_text(candidate)
                if API_HINT_RE.search(candidate):
                    findings.append({"script_url": url, "candidate": candidate[:500]})
        except Exception as exc:  # noqa: BLE001
            state.errors.append(f"inspect_javascript: {url}: {exc}")

    write_csv(state.output_dir / "discovered_endpoints.csv", findings)


def iter_dict_records(value: Any, path: str = "$", depth: int = 0) -> Iterable[tuple[str, dict[str, Any]]]:
    if depth > 12:
        return
    if isinstance(value, dict):
        yield path, value
        for key, child in value.items():
            yield from iter_dict_records(child, f"{path}.{key}", depth + 1)
    elif isinstance(value, list):
        for idx, child in enumerate(value):
            yield from iter_dict_records(child, f"{path}[{idx}]", depth + 1)


def scalarize(record: dict[str, Any]) -> dict[str, str]:
    output: dict[str, str] = {}
    for key, value in record.items():
        if isinstance(value, (str, int, float, bool)) or value is None:
            output[str(key)] = clean_text(value)
        elif isinstance(value, list) and all(isinstance(x, (str, int, float, bool)) or x is None for x in value):
            output[str(key)] = " | ".join(clean_text(x) for x in value)
    return output


def classify_record(record: dict[str, str]) -> set[str]:
    keys = " ".join(record.keys()).lower()
    values = " ".join(record.values()).lower()
    combined = keys + " " + values
    classes: set[str] = set()
    if re.search(r"operator|company|kompani|subjekt|shoqeri", combined):
        classes.add("operators")
    if re.search(r"agency|agjenci", combined):
        classes.add("agencies")
    if re.search(r"station|terminal|stacion", combined):
        classes.add("stations")
    if re.search(r"route|line|linj|departure|arrival|nisje|mberrit|mbërrit|schedule|orar", combined):
        classes.add("routes")
    if re.search(r"price|fare|tariff|cmim|çmim|amount|currency", combined) or PRICE_RE.search(values):
        classes.add("prices")
    if PHONE_RE.search(values) or EMAIL_RE.search(values):
        classes.add("contacts")
    return classes


def extract_structured(state: State) -> dict[str, list[dict[str, str]]]:
    groups: dict[str, list[dict[str, str]]] = defaultdict(list)
    dedupe: dict[str, set[str]] = defaultdict(set)

    for item in state.json_payloads:
        source_url = item["source_url"]
        for path, raw in iter_dict_records(item["payload"]):
            record = scalarize(raw)
            if not record:
                continue
            record["_source_url"] = source_url
            record["_json_path"] = path
            fingerprint = hashlib.sha256(json.dumps(record, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
            groups["raw_records"].append(record)
            for category in classify_record(record):
                if fingerprint not in dedupe[category]:
                    dedupe[category].add(fingerprint)
                    groups[category].append(record)

    # Add DOM blocks to contacts and route/price groups.
    for block in state.blocks:
        record = {k: clean_text(v) for k, v in block.items()}
        if record.get("phones") or record.get("emails"):
            groups["contacts"].append(record)
        if record.get("prices") or record.get("times"):
            groups["routes"].append(record)
            if record.get("prices"):
                groups["prices"].append(record)
        if "sh.p.k" in record.get("text", "").lower() or "operator" in record.get("text", "").lower():
            groups["operators"].append(record)

    return groups


def union_headers(rows: list[dict[str, Any]]) -> list[str]:
    preferred = [
        "name", "companyName", "operatorName", "agencyName", "phone", "phoneNumber",
        "email", "price", "currency", "departure", "arrival", "departureTime",
        "arrivalTime", "origin", "destination", "city", "address", "text", "phones",
        "emails", "prices", "times", "_source_url", "_json_path", "page_url",
    ]
    keys: list[str] = []
    seen: set[str] = set()
    all_keys = {str(k) for row in rows for k in row.keys()}
    for key in preferred + sorted(all_keys):
        if key in all_keys and key not in seen:
            seen.add(key)
            keys.append(key)
    return keys


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = union_headers(rows) if rows else ["no_data"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        if rows:
            writer.writerows(rows)


def style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions
    for col_idx, column in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in list(column)[:500]:
            max_len = max(max_len, len(clean_text(cell.value)))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 55)


def add_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=name[:31])
    headers = union_headers(rows) if rows else ["no_data"]
    ws.append(headers)
    for row in rows:
        ws.append([clean_text(row.get(h, "")) for h in headers])
    style_sheet(ws)


def export_all(state: State, groups: dict[str, list[dict[str, str]]]) -> None:
    for category, rows in groups.items():
        write_csv(state.output_dir / f"{category}.csv", rows)
    write_csv(state.output_dir / "network_responses.csv", state.responses)
    write_csv(state.output_dir / "visited_pages.csv", state.visited)
    write_csv(state.output_dir / "dom_blocks.csv", state.blocks)
    write_csv(state.output_dir / "errors.csv", [{"error": e} for e in state.errors])

    summary = [
        {"metric": "Scraped at (UTC)", "value": datetime.now(timezone.utc).isoformat()},
        {"metric": "Visited pages", "value": len(state.visited)},
        {"metric": "Captured JSON responses", "value": len(state.responses)},
        {"metric": "Parsed JSON payloads", "value": len(state.json_payloads)},
        {"metric": "Operator-like records", "value": len(groups.get("operators", []))},
        {"metric": "Agency-like records", "value": len(groups.get("agencies", []))},
        {"metric": "Station-like records", "value": len(groups.get("stations", []))},
        {"metric": "Route/schedule-like records", "value": len(groups.get("routes", []))},
        {"metric": "Price-like records", "value": len(groups.get("prices", []))},
        {"metric": "Contact-like records", "value": len(groups.get("contacts", []))},
        {"metric": "Errors", "value": len(state.errors)},
    ]

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Summary", summary)
    add_sheet(wb, "Operators", groups.get("operators", []))
    add_sheet(wb, "Agencies", groups.get("agencies", []))
    add_sheet(wb, "Stations", groups.get("stations", []))
    add_sheet(wb, "Routes_Schedules", groups.get("routes", []))
    add_sheet(wb, "Prices", groups.get("prices", []))
    add_sheet(wb, "Contacts", groups.get("contacts", []))
    add_sheet(wb, "Network_Responses", state.responses)
    add_sheet(wb, "Visited_Pages", state.visited)
    add_sheet(wb, "Raw_Records", groups.get("raw_records", [])[:50000])
    add_sheet(wb, "Errors", [{"error": e} for e in state.errors])
    wb.save(state.output_dir / "etransport_data.xlsx")

    manifest = {
        "base_url": BASE_URL,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "counts": {k: len(v) for k, v in groups.items()},
        "responses": len(state.responses),
        "visited_pages": len(state.visited),
        "errors": state.errors,
    }
    (state.output_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
    )


async def run(args: argparse.Namespace) -> None:
    output_dir = Path(args.output).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    state = State(output_dir=output_dir)

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=not args.show_browser)
        context = await browser.new_context(
            locale="sq-AL",
            timezone_id="Europe/Tirane",
            viewport={"width": 1440, "height": 1100},
            user_agent=(
                "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/150.0.0.0 Safari/537.36"
            ),
        )

        queue = [canonical_url(urljoin(BASE_URL, path)) for path in START_PATHS]
        visited_urls: set[str] = set()

        while queue and len(visited_urls) < args.max_pages:
            url = canonical_url(queue.pop(0))
            if url in visited_urls or not same_origin(url):
                continue
            visited_urls.add(url)
            print(f"[{len(visited_urls)}/{args.max_pages}] {url}")
            await visit(context, url, state)

            if args.crawl_links:
                for link in sorted(state.links):
                    if link not in visited_urls and link not in queue:
                        path = urlparse(link).path.lower()
                        if any(token in path for token in ["rrjetilinjave", "information", "search"]):
                            queue.append(link)

        await inspect_javascript(context, state)
        await context.close()
        await browser.close()

    groups = extract_structured(state)
    export_all(state, groups)
    print(f"\nDone. Results: {output_dir}")
    print(f"Excel: {output_dir / 'etransport_data.xlsx'}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Scrape public eTransport data")
    parser.add_argument("--output", default="etransport_output", help="Output folder")
    parser.add_argument("--max-pages", type=int, default=40, help="Maximum public pages to visit")
    parser.add_argument("--show-browser", action="store_true", help="Show Chromium while scraping")
    parser.add_argument("--no-crawl-links", dest="crawl_links", action="store_false", help="Do not crawl discovered internal links")
    parser.set_defaults(crawl_links=True)
    return parser.parse_args()


if __name__ == "__main__":
    try:
        asyncio.run(run(parse_args()))
    except KeyboardInterrupt:
        sys.exit(130)
